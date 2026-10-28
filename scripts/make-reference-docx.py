# -*- coding: utf-8 -*-
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC = sys.argv[1]
DST = sys.argv[2]

doc = Document(SRC)
styles = doc.styles


def set_font(style, ascii_font, east_font, size=None, bold=None, color=None):
    f = style.font
    f.name = ascii_font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = RGBColor(*color)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_font)
    rFonts.set(qn("w:cs"), ascii_font)


def shade(style_or_el, fill):
    el = style_or_el.element if hasattr(style_or_el, "element") else style_or_el
    pPr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


# 基本样式
set_font(styles["Normal"], "Times New Roman", "宋体", 10.5)
styles["Normal"].paragraph_format.line_spacing = 1.45
styles["Normal"].paragraph_format.space_after = Pt(2)

for name in ("Body Text", "First Paragraph", "Compact"):
    try:
        st = styles[name]
    except KeyError:
        continue
    set_font(st, "Times New Roman", "宋体", 10.5)
    pf = st.paragraph_format
    pf.line_spacing = 1.45
    pf.space_after = Pt(3)
    pf.first_line_indent = Pt(21)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# 标题
h1 = styles["Heading 1"]
set_font(h1, "Times New Roman", "黑体", 18, bold=False, color=(0, 0, 0))
h1.paragraph_format.space_before = Pt(18)
h1.paragraph_format.space_after = Pt(8)
h1.paragraph_format.line_spacing = 1.2
h1.paragraph_format.keep_with_next = True
h1.paragraph_format.first_line_indent = Pt(0)

h2 = styles["Heading 2"]
set_font(h2, "Times New Roman", "黑体", 14, bold=False, color=(0, 0, 0))
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.line_spacing = 1.2
h2.paragraph_format.keep_with_next = True
h2.paragraph_format.first_line_indent = Pt(0)

h3 = styles["Heading 3"]
set_font(h3, "Times New Roman", "黑体", 12, bold=True, color=(0, 0, 0))
h3.paragraph_format.space_before = Pt(8)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True
h3.paragraph_format.first_line_indent = Pt(0)

try:
    t = styles["Title"]
    set_font(t, "Times New Roman", "黑体", 26, bold=True, color=(0, 0, 0))
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(6)
except KeyError:
    pass

for name, size, font in (("Subtitle", 15, "楷体"), ("Author", 11, "宋体"), ("Date", 11, "宋体")):
    try:
        st = styles[name]
    except KeyError:
        continue
    set_font(st, "Times New Roman", font, size, bold=False, color=(0x40, 0x40, 0x40))
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(4)

# 代码样式（pandoc 使用 Source Code / Verbatim Char）
try:
    sc = styles["Source Code"]
except KeyError:
    sc = styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
set_font(sc, "Consolas", "等线", 8.5, bold=False, color=(0x1F, 0x1F, 0x1F))
sc.paragraph_format.line_spacing = 1.05
sc.paragraph_format.space_before = Pt(3)
sc.paragraph_format.space_after = Pt(6)
sc.paragraph_format.first_line_indent = Pt(0)
sc.paragraph_format.left_indent = Cm(0.3)
shade(sc, "F4F5F7")

try:
    vc = styles["Verbatim Char"]
except KeyError:
    vc = styles.add_style("Verbatim Char", WD_STYLE_TYPE.CHARACTER)
set_font(vc, "Consolas", "等线", 9)

# 表格样式
table_style = None
for s in styles:
    if s.name == "Table" and s.type == WD_STYLE_TYPE.TABLE:
        table_style = s
        break
if table_style is None:
    table_style = styles.add_style("Table", WD_STYLE_TYPE.TABLE)
table_style.font.size = Pt(9)
table_style.font.name = "Times New Roman"
rPr = table_style.element.get_or_add_rPr()
rFonts = rPr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = OxmlElement("w:rFonts")
    rPr.insert(0, rFonts)
rFonts.set(qn("w:ascii"), "Times New Roman")
rFonts.set(qn("w:hAnsi"), "Times New Roman")
rFonts.set(qn("w:eastAsia"), "宋体")
tblPr = table_style.element.find(qn("w:tblPr"))
if tblPr is None:
    tblPr = OxmlElement("w:tblPr")
    table_style.element.append(tblPr)
for border in ("top", "left", "bottom", "right", "insideH", "insideV"):
    b = OxmlElement("w:" + border)
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), "808080")
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    borders.append(b)

# 页面与页脚
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.6)

footer = sec.footer
footer.is_linked_to_previous = False
p = footer.paragraphs[0]
p.text = ""
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("维修百事通 · 提示词工程规范 v1.0    第 ")
run.font.size = Pt(9)
run.font.name = "Times New Roman"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
r2 = p.add_run()
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
r2._r.append(fld1); r2._r.append(instr); r2._r.append(fld2)
r2.font.size = Pt(9)
run3 = p.add_run(" 页")
run3.font.size = Pt(9)
run3.font.name = "Times New Roman"
run3._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
run3.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

doc.save(DST)
print("saved", DST)
